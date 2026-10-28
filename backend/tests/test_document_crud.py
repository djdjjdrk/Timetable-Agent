"""文档接口验证：隔离数据库与文件目录，不修改实际课程资料。"""

from pathlib import Path
from io import BytesIO
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from starlette.exceptions import HTTPException

from app.api.document import router
from app.common.exceptions import BusinessException, bussiness_excpetion_hadler, http_excpetion_hadler
from app.config import settings
from app.database import Base, get_db
from app.models import User
from app.services import document_service
from app.utils.jwt import create_access_token


class DocumentCrudTests(unittest.TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.directory = Path(directory.name)
        for patcher in (
            patch.object(document_service, 'KB_DIR', self.directory),
            patch.object(settings, 'JWT_SECRET_KEY', 'document-tests-only-secret-key-32-characters'),
            # CRUD 测试验证文件与事务；向量写入隔离，避免触碰实际 Chroma。
            patch.object(document_service, 'ingest_kb_file',
                         side_effect=lambda path, **_: document_service.extract_document_text(path)),
            patch.object(document_service, 'delete_by_doc_id'),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)
        engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        self.addCleanup(engine.dispose)
        Base.metadata.create_all(engine)
        self.factory = sessionmaker(bind=engine)
        with self.factory() as db:
            for user_id, role in [(1, 'admin'), (2, 'student')]:
                db.add(User(id=user_id, username=role, name=role, password='unused', role=role, status=1))
            db.commit()

        def get_test_db():
            with self.factory() as db:
                yield db

        app = FastAPI()
        app.add_exception_handler(BusinessException, bussiness_excpetion_hadler)
        app.add_exception_handler(HTTPException, http_excpetion_hadler)
        app.include_router(router, prefix='/api')
        app.dependency_overrides[get_db] = get_test_db
        self.client = TestClient(app)
        self.addCleanup(self.client.close)

    def request(self, method, suffix='', user_id=1, **kwargs):
        headers = {} if user_id is None else {'Authorization': f'Bearer {create_access_token(user_id)}'}
        return self.client.request(method, '/api/documents' + suffix, headers=headers, **kwargs)

    def test_four_formats_upload_preview_delete_and_duplicate(self):
        doc = Document()
        doc.add_paragraph('数据库课程资料。' * 20)
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'description'
        table.cell(0, 1).text = '允许为空'
        docx_bytes = BytesIO()
        doc.save(docx_bytes)
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 50 700 Td (' + b'TCP reliable transport course. ' * 5 + b') Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        pdf_bytes = BytesIO()
        writer.write(pdf_bytes)
        files = [
            ('files', ('math.txt', ('导数课程笔记。' * 20).encode('utf-8'))),
            ('files', ('structure.md', ('# 栈与队列\n后进先出与先进先出。\n' * 10).encode('utf-8'))),
            ('files', ('database.docx', docx_bytes.getvalue())),
            ('files', ('network.pdf', pdf_bytes.getvalue())),
        ]
        response = self.request('POST', files=files).json()
        self.assertEqual(response['code'], 200, response)
        self.assertEqual(len(response['data']), 4)
        self.assertEqual(len(list(self.directory.iterdir())), 4)
        listed = self.request('GET').json()['data']
        self.assertEqual(len(listed), 4)
        for row in listed:
            preview = self.request('GET', f"/{row['id']}/preview").json()
            self.assertEqual(preview['code'], 200, preview)
            data = preview['data']
            self.assertEqual(data['char_count'], len(data['text']))
            self.assertGreater(data['char_count'], 100)
            if row['ext'] == '.docx':
                self.assertIn('description', data['text'])
                self.assertIn('允许为空', data['text'])
            if row['ext'] == '.pdf':
                self.assertIn('reliable transport', data['text'])
        duplicate = self.request('POST', files=[files[0]]).json()
        self.assertEqual(duplicate['code'], 400)
        self.assertIn('同名', duplicate['message'])
        self.assertEqual(len(list(self.directory.iterdir())), 4)
        row = listed[0]
        self.assertEqual(self.request('DELETE', f"/{row['id']}").json()['code'], 200)
        self.assertFalse((self.directory / row['rel_path']).exists())
        self.assertEqual(self.request('GET', f"/{row['id']}/preview").json()['code'], 404)

    def test_batch_rollback_and_blank_preview(self):
        for files in [
            [('files', ('valid.txt', b'course notes')), ('files', ('empty.txt', b''))],
            [('files', ('same.txt', b'first')), ('files', ('same.txt', b'second'))],
        ]:
            self.assertEqual(self.request('POST', files=files).json()['code'], 400)
            self.assertEqual(self.request('GET').json()['data'], [])
            self.assertEqual(list(self.directory.iterdir()), [])
        blank = self.request('POST', files=[('files', ('blank.txt', b'   \n'))]).json()
        self.assertEqual(blank['code'], 400, blank)
        self.assertEqual(self.request('GET').json()['data'], [])
        self.assertEqual(list(self.directory.iterdir()), [])
        # 新上传会拒绝空白资料；旧资料文件变空后，预览仍应正确报错。
        row = self.request('POST', files=[('files', ('legacy.txt', b'course notes'))]).json()['data'][0]
        (self.directory / row['rel_path']).write_text('   \n', encoding='utf-8')
        self.assertEqual(self.request('GET', f"/{row['id']}/preview").json()['code'], 400)

    def test_all_four_endpoints_require_admin(self):
        for method, suffix, kwargs in [
            ('GET', '', {}),
            ('POST', '', {'files': [('files', ('course.txt', b'notes'))]}),
            ('GET', '/1/preview', {}),
            ('DELETE', '/1', {}),
        ]:
            student = self.request(method, suffix, user_id=2, **kwargs)
            self.assertEqual(student.status_code, 200)
            self.assertEqual(student.json()['code'], 403)
            anonymous = self.request(method, suffix, user_id=None, **kwargs)
            self.assertEqual(anonymous.status_code, 401)


if __name__ == '__main__':
    unittest.main()
